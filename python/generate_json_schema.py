from docx import Document
import json
import os

# --------------------> 全局变量  <--------------------
# 字段类型字典
schema_type_dic = {"字符串":"string", "整型": "integer", "对象": "object", "数组": "array", "布尔": "boolean", "空": "null", "浮点型": "number", 
"JSON对象": "object", "JSON数组": "array", "JSON字符串数组": "array", "JSON整型数组": "array"}




def write_json_schema_file(num, action_name, the_json):
    ''' 写入 json schema 到文件 '''
    if not os.path.exists('json_schema_collections'):
        os.mkdir('json_schema_collections')
    with open("json_schema_collections/"+str(num)+"_"+action_name+".json","w",encoding='utf-8') as f:
        json.dump(the_json,f,ensure_ascii=False)


def parse_lines(i, table, level, parent_node, father_is_object_array):
    ''' 处理一行的数据 '''
    print("---------------------------------------------------------------------------------------------------------------")
    print(i) 

    if i >= len(table.rows):
        print("return 1")
        return i        
    cur_level = table.cell(i, 0).text.count('+')
    print("### cur_level:" + str(cur_level) +"    level:"+str(level) )
    if cur_level != level or table.cell(i, 0).text == "":
        print("return 2")
        return i-1
    
    # 获得一行中的所有字段的值
    name = table.cell(i, 0).text
    real_name = name.replace('+', '')
    definition = table.cell(i, 1).text
    required = table.cell(i, 2).text
    restriction = table.cell(i, 3).text
    the_type = table.cell(i, 4).text.replace(' ', '')
    description = table.cell(i, 5).text

    # 层级。值为 0 时代表顶层，值为 1 时代表子第一层
    
    print(real_name + " - " + str(level))
    # 得到父级 json 节点
    cur_parent_node =  parent_node


    # 设置节点
    msg = ("取值范围：" + restriction +  "；" if restriction != "" else "") + (description if  description != "" else "")
    cur_parent_node['properties'][real_name] = { "type": schema_type_dic[the_type], "title": definition, "description":  msg}


    if the_type == "JSON对象":
        # ！！！ 设置了父级节点
        # print("--------> 2 <-------")
        cur_parent_node['properties'][real_name]['properties'] = {}
        cur_parent_node['properties'][real_name]['required'] = []
        i = parse_lines(i+1, table, level+1, cur_parent_node['properties'][real_name], 0)
    elif the_type == "JSON数组":
        # 这里其实是一个 json 对象数组
        # ！！！ 设置了父级节点
        # cur_parent_node['properties'][real_name]['properties'] = {}
        # cur_parent_node['properties'][real_name]['required'] = []
        cur_parent_node['properties'][real_name]['items'] = {"type": "object"}
        cur_parent_node['properties'][real_name]['items']['properties'] = {}
        cur_parent_node['properties'][real_name]['items']['required'] = []
        i = parse_lines(i+1, table, level+1, cur_parent_node['properties'][real_name]['items'], 1)
    elif the_type == "JSON字符串数组":
        # OK
        cur_parent_node['properties'][real_name]['items'] = {"type": "string"}
    elif the_type == "JSON整型数组":
        # OK
        cur_parent_node['properties'][real_name]['items'] = {"type": "number"}

    if required == "Y":
        required_list = cur_parent_node['required']
        required_list.append(real_name)
        cur_parent_node['required'] = required_list
        # print("##################################")
        # print(cur_parent_node['required'])

    # 递归。进行下一行的解析
    i = parse_lines(i+1, table, level, parent_node, father_is_object_array)
    return i





def parse_one_table(table, num):
    ''' 解析一个表格 '''
    # 名称	    定义	            必需	长度/范围	    类型	        备注
    # Action	固定“addPerson”	    Y		                字符串	
    
    # 生成 json schema 的根节点
    the_json = json.loads("{}")
    the_json['type'] = 'object'
    the_json['properties'] = {}
    the_json['required'] = []

    # 解析表格中的每一行
    # for i in range(1, len(table.rows)):
    #     name = table.cell(i, 0).text
    #     if name != "":
    #         level = name.count('+')

    # 开启递归，解析表格
    parse_lines(1, table, 0, the_json, 0)

    # 写入 json schema 到文件是
    print(the_json)
    action_name = table.cell(1, 1).text
    write_json_schema_file(num, action_name, the_json)


def parse_tables(tables):
    '''解析 tables 表格'''
    count = 0
    for i in range(len(tables)):
        if tables[i].cell(1,0).text == "Action" and len(tables[i].columns) == 6:
            count += 1
            parse_one_table(tables[i], count)
            # break


def load_docx_tables(path):
    '''加载 docx 文件中的表格'''
    document = Document(path)
    return document.tables


def main():
    '''程序入口'''
    print("Start parse docx")
    tables = load_docx_tables("D:/0-DesktopData/python-mqtt-client/docx/mqtt.docx")
    parse_tables(tables)


main()